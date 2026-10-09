"""Fixed-template reconciliation. Imports create drafts, never confirmed costs."""
from collections import Counter
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP
from io import BytesIO
import hashlib
import json
import re
import uuid
from zipfile import BadZipFile, ZipFile

from openpyxl import load_workbook
from sqlalchemy import func, or_
from sqlalchemy.orm import aliased

from Auth import get_payload_roles
from data.DB import CargoTable, OrderTable
from data.finance.ledger import FinanceDocument, FinanceDocumentLine, FinanceVendor
from data.supplier_reconciliation import SupplierReconciliationBatch, SupplierReconciliationRow
from services import finance_ledger_service as ledger

MAX_FILE_BYTES = 10 * 1024 * 1024
MAX_ROWS = 5000
MIN_WEIGHT_DIFFERENCE = Decimal("-0.5")
MIN_ORDER_WEIGHT_DIFFERENCE = Decimal("-5")
MAX_WEIGHT_DIFFERENCE = Decimal("0")
WEIGHT = Decimal("0.001")
HEADERS = ("原单号", "转单号", "重量(kg)", "最终费用(元)", "费用说明")
REQUIRED_HEADERS = {"原单号", "重量(kg)", "最终费用(元)"}


def _actor(actor):
    if not set(get_payload_roles(actor or {})).intersection({"admin", "clerk", "finance"}):
        raise PermissionError("仅文员、财务或管理员可以处理供应商对账明细")
    actor_id = (actor or {}).get("user_id") or (actor or {}).get("id")
    if not actor_id:
        raise PermissionError("无法识别操作人")
    return int(actor_id), str(actor.get("name") or actor.get("username") or actor_id)


def _json(value):
    return json.dumps(value, ensure_ascii=False, sort_keys=True, separators=(",", ":"))


def _issue(code, message):
    return {"code": code, "message": message}


def _weight_issue(difference):
    """Difference is uploaded weight minus the unrounded warehouse weight."""
    if difference > MAX_WEIGHT_DIFFERENCE:
        return _issue("weight_mismatch", f"上传重量大于系统实重（高出 {format(difference, 'f')} kg）")
    if difference < MIN_WEIGHT_DIFFERENCE:
        return _issue("weight_mismatch", f"上传重量比系统实重少 {format(-difference, 'f')} kg，超过 0.5 kg")
    return None


def _order_weight_issue(difference):
    if difference >= MAX_WEIGHT_DIFFERENCE:
        return _issue("weight_mismatch", "订单级上传重量必须小于系统订单总实重")
    if difference < MIN_ORDER_WEIGHT_DIFFERENCE:
        return _issue("weight_mismatch", f"上传重量比系统订单总实重少 {format(-difference, 'f')} kg，超过 5 kg")
    return None


def _cell_number(cell, field, issues, *, required=False):
    if cell is None or cell.value is None or str(cell.value).strip() == "":
        if required:
            issues.append(_issue("missing_original_number", "原单号不能为空"))
        return ""
    value = cell.value
    if cell.data_type in {"f", "e"} or isinstance(value, bool):
        issues.append(_issue("invalid_number", f"{field}必须粘贴为文本，不能使用公式或错误值"))
        return ""
    if isinstance(value, (int, float, Decimal)):
        try:
            decimal = Decimal(str(value))
            if not decimal.is_finite() or decimal < 0 or decimal != decimal.to_integral_value():
                raise ValueError()
            text = str(int(decimal))
            if len(text) > 15:
                issues.append(_issue("unsafe_excel_number", f"{field}超过 Excel 数字安全精度，请重新从原账单粘贴为文本"))
                return ""
            if re.fullmatch(r"0+", cell.number_format or ""):
                text = text.zfill(len(cell.number_format))
        except (InvalidOperation, ValueError, OverflowError):
            issues.append(_issue("invalid_number", f"{field}必须为有效文本号码"))
            return ""
    elif isinstance(value, str):
        text = value.strip()
        if re.fullmatch(r"[+-]?\d+(?:\.\d+)?[eE][+-]?\d+", text):
            issues.append(_issue("unsafe_excel_number", f"{field}不能使用科学计数法，请粘贴完整文本号码"))
            return ""
    else:
        issues.append(_issue("invalid_number", f"{field}必须为文本号码"))
        return ""
    if len(text) > 256:
        issues.append(_issue("invalid_number", f"{field}不能超过 256 字"))
        return ""
    return text


def _cell_decimal(cell, field, precision, issues):
    value = cell.value if cell is not None else None
    try:
        if value is None or isinstance(value, bool) or cell.data_type in {"f", "e"}:
            raise ValueError()
        decimal = Decimal(str(value).strip())
        if not decimal.is_finite() or abs(decimal) > Decimal("999999999999.99"):
            raise ValueError()
        if decimal != decimal.quantize(precision):
            issues.append(_issue("invalid_precision", f"{field}最多保留 {'3' if precision == WEIGHT else '2'} 位小数"))
            return None
        if decimal <= 0:
            issues.append(_issue("nonpositive_weight" if precision == WEIGHT else "nonpositive_fee",
                "重量必须大于 0" if precision == WEIGHT else "最终费用必须大于 0；负数扣款请由财务另行处理供应商贷项"))
        return decimal.quantize(precision)
    except (InvalidOperation, TypeError, ValueError):
        issues.append(_issue("invalid_weight" if precision == WEIGHT else "invalid_fee", f"{field}必须为有效数字，公式请先粘贴为值"))
        return None


def parse_workbook(content):
    """Validate the file schema, preserving row errors as reconciliation issues."""
    if not isinstance(content, bytes) or not content or len(content) > MAX_FILE_BYTES:
        raise ValueError("请上传不超过 10 MB 的 .xlsx 文件")
    try:
        with ZipFile(BytesIO(content)) as archive:
            if len(archive.infolist()) > 1000 or sum(info.file_size for info in archive.infolist()) > 100 * 1024 * 1024:
                raise ValueError("Excel 解压后大小超出限制")
            if "xl/workbook.xml" not in archive.namelist():
                raise ValueError("文件不是有效的 .xlsx 工作簿")
        workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    except (BadZipFile, KeyError, OSError) as exc:
        raise ValueError("无法读取 Excel，请使用下载的统一模板") from exc
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("无法读取 Excel，请使用下载的统一模板") from exc
    try:
        if "对账明细" not in workbook.sheetnames:
            raise ValueError("缺少工作表“对账明细”，请使用统一模板")
        sheet = workbook["对账明细"]
        if sheet.max_row is None or sheet.max_row > MAX_ROWS + 1 or (sheet.max_column or 0) > 20:
            raise ValueError("对账明细最多支持 5000 行，且不能包含模板之外的大量列")
        header_cells = next(sheet.iter_rows(min_row=1, max_row=1), ())
        positions = {}
        for index, cell in enumerate(header_cells):
            if cell.value is None or str(cell.value).strip() == "":
                continue
            header = str(cell.value).strip()
            if cell.data_type == "f" or header not in HEADERS or header in positions:
                raise ValueError("标题行必须使用统一模板列名，不能重复或增加其它列")
            positions[header] = index
        if not REQUIRED_HEADERS.issubset(positions):
            raise ValueError("标题行缺少原单号、重量(kg)或最终费用(元)")
        rows = []
        for row_number, cells in enumerate(sheet.iter_rows(min_row=2), 2):
            if all(cell.value is None or str(cell.value).strip() == "" for cell in cells):
                continue
            get_cell = lambda header: cells[positions[header]] if header in positions else None
            issues = []
            raw = {header: {"value": str(get_cell(header).value)[:4096] if get_cell(header).value is not None else None,
                           "type": get_cell(header).data_type, "number_format": get_cell(header).number_format}
                   for header in positions}
            original = _cell_number(get_cell("原单号"), "原单号", issues, required=True)
            waybill = _cell_number(get_cell("转单号"), "转单号", issues)
            weight = _cell_decimal(get_cell("重量(kg)"), "重量(kg)", WEIGHT, issues)
            fee = _cell_decimal(get_cell("最终费用(元)"), "最终费用(元)", ledger.MONEY, issues)
            description_cell = get_cell("费用说明")
            description = str(description_cell.value or "").strip() if description_cell else ""
            if description_cell and description_cell.data_type in {"f", "e"}:
                issues.append(_issue("invalid_description", "费用说明不能使用公式或错误值"))
            if len(description) > 512:
                issues.append(_issue("invalid_description", "费用说明不能超过 512 字"))
            rows.append({"source_row": row_number, "raw_json": _json(raw), "original_number": original,
                "waybill_number": waybill, "weight": weight, "final_fee": fee, "description": description[:512],
                "parse_issues_json": _json(issues)})
        if not rows:
            raise ValueError("对账明细没有数据行")
        return rows
    finally:
        workbook.close()


def _row_dict(row):
    result = {field: getattr(row, field) for field in ("id", "source_row", "original_number", "waybill_number", "description",
        "status", "cargo_uid", "order_uid", "order_number", "system_waybill", "finance_document_id", "finance_document_line_id")}
    for field, precision in (("weight", WEIGHT), ("system_weight", WEIGHT), ("weight_difference", WEIGHT), ("final_fee", ledger.MONEY)):
        value = getattr(row, field)
        result[field] = ledger._str(value, precision) if value is not None else None
    result["issues"] = json.loads(row.issues_json)
    result["match_type"] = "cargo" if row.cargo_uid else "order" if row.order_uid else None
    result["system_weight_scope"] = "order_total" if result["match_type"] == "order" else "cargo" if row.cargo_uid else None
    result["raw_values"] = {header: cell.get("value") for header, cell in json.loads(row.raw_json).items()}
    return result


def _documents(session, batch_id):
    ids = [value[0] for value in session.query(SupplierReconciliationRow.finance_document_id).filter_by(batch_id=batch_id)
           .filter(SupplierReconciliationRow.finance_document_id.isnot(None)).distinct()]
    return [{"id": doc.id, "number": doc.number, "status": doc.status, "total_amount": ledger._str(doc.total_amount), "version": doc.version}
            for doc in session.query(FinanceDocument).filter(FinanceDocument.id.in_(ids)).order_by(FinanceDocument.id)] if ids else []


def _batch_dict(session, batch):
    result = {field: getattr(batch, field) for field in ("id", "number", "vendor_id", "vendor_name", "period", "filename",
        "status", "version", "created_at", "updated_at", "created_by_name")}
    result["summary"] = {"total_rows": batch.total_rows, "passed_rows": batch.passed_rows, "exception_rows": batch.exception_rows,
        "passed_amount": ledger._str(batch.passed_amount), "total_valid_amount": ledger._str(batch.total_valid_amount)}
    result["documents"] = _documents(session, batch.id)
    result["weight_tolerance"] = ledger._str(-MIN_WEIGHT_DIFFERENCE, WEIGHT)
    result["weight_difference_range"] = {"min": ledger._str(MIN_WEIGHT_DIFFERENCE, WEIGHT), "max": ledger._str(MAX_WEIGHT_DIFFERENCE, WEIGHT)}
    result["order_weight_difference_range"] = {"min": ledger._str(MIN_ORDER_WEIGHT_DIFFERENCE, WEIGHT),
        "max": ledger._str(MAX_WEIGHT_DIFFERENCE, WEIGHT), "min_inclusive": True, "max_inclusive": False}
    result["notice"] = "先匹配包裹原单号，无包裹候选时匹配订单号。包裹重量差允许 -0.5 至 0 kg（含边界）；订单按未归档包裹总实重核对，重量差允许 -5 至 0 kg（不含 0），上传重量必须小于订单总实重。最终费用按上传金额生成应付草稿，未校验供应商报价。异常行不入账，草稿生效前不计入实际成本。"
    return result


def _detail(session, batch, *, page=1, page_size=50, row_status=""):
    page, page_size = max(int(page or 1), 1), min(max(int(page_size or 50), 1), 200)
    query = session.query(SupplierReconciliationRow).filter_by(batch_id=batch.id)
    if row_status:
        if row_status not in {"passed", "exception"}:
            raise ValueError("明细状态无效")
        query = query.filter_by(status=row_status)
    total = query.count()
    rows = query.order_by(SupplierReconciliationRow.source_row).offset((page - 1) * page_size).limit(page_size).all()
    return {"batch": _batch_dict(session, batch), "rows": {"items": [_row_dict(row) for row in rows],
        "total": total, "page": page, "page_size": page_size}}


def _snapshot(cargo, order):
    raw_weight = Decimal(str(cargo.weight or 0))
    if not raw_weight.is_finite() or raw_weight <= 0:
        raise ValueError("系统重量无效")
    return {"cargo_uid": cargo.uid, "order_uid": cargo.target_order_uid,
        "order_number": order.order_id or order.uid, "original_number": str(cargo.transfer_id or "").strip(),
        "waybill_number": str(cargo.waybill_id or "").strip(), "weight": ledger._str(cargo.weight, WEIGHT),
        "weight_raw": str(raw_weight),
        "cargo_archive": bool(cargo.archive), "order_archive": bool(order.archive)}


def _order_snapshot(order, cargos):
    """Real weight uses the same active-cargo scope as current_order_weight.

    Sum the original decimals before display rounding; total_weight and billing
    weights are not warehouse weighing evidence. Record membership as well as
    each weight so changing boxes without changing the total requires review.
    """
    members, total = [], Decimal("0")
    for cargo in sorted(cargos, key=lambda value: value.uid):
        raw = Decimal(str(cargo.weight))
        if not raw.is_finite() or raw < 0:
            raise ValueError("系统订单重量无效")
        total += raw
        members.append({"cargo_uid": cargo.uid, "weight_raw": str(raw),
            "waybill_number": str(cargo.waybill_id or "").strip()})
    if not members or not total.is_finite() or total <= 0:
        raise ValueError("系统订单重量无效")
    return {"match_type": "order", "cargo_uid": None, "order_uid": order.uid,
        "order_number": order.order_id or order.uid, "original_number": str(order.order_id or "").strip(),
        "waybill_number": None, "weight": ledger._str(total, WEIGHT), "weight_raw": str(total),
        "order_archive": bool(order.archive), "cargos": members}


def _matching_sources(session, originals):
    original_set = set(originals)
    originals = sorted(original_set)
    candidate_uids, candidate_orders = set(), set()
    cargo_numbers = set()
    for start in range(0, len(originals), 500):
        for uid, order_uid, number in session.query(CargoTable.uid, CargoTable.target_order_uid, CargoTable.transfer_id).filter(
                func.trim(CargoTable.transfer_id).in_(originals[start:start + 500])):
            # MySQL collations may be case-insensitive; identifiers still match exactly.
            if str(number or "").strip() not in original_set:
                continue
            candidate_uids.add(uid)
            cargo_numbers.add(str(number or "").strip())
            if order_uid:
                candidate_orders.add(order_uid)
    fallback_numbers = [number for number in originals if number not in cargo_numbers]
    fallback_set = set(fallback_numbers)
    fallback_orders = set()
    for start in range(0, len(fallback_numbers), 500):
        for uid, number in session.query(OrderTable.uid, OrderTable.order_id).filter(
                func.trim(OrderTable.order_id).in_(fallback_numbers[start:start + 500])):
            if str(number or "").strip() in fallback_set:
                candidate_orders.add(uid)
                fallback_orders.add(uid)
    # Business mutations lock orders before cargo; keep that same order here.
    orders = {order.uid: order for order in session.query(OrderTable).filter(OrderTable.uid.in_(sorted(candidate_orders)))
              .order_by(OrderTable.uid).populate_existing().with_for_update().all()}
    candidates, order_candidates, order_cargos = {}, {}, {}
    for order in orders.values():
        if order.uid in fallback_orders:
            order_candidates.setdefault(str(order.order_id or "").strip(), []).append(order)
    for cargo in session.query(CargoTable).filter(or_(CargoTable.uid.in_(sorted(candidate_uids)),
            (CargoTable.target_order_uid.in_(sorted(fallback_orders))) & (CargoTable.archive.is_(False)))).order_by(
            CargoTable.uid).populate_existing().with_for_update().all():
        if cargo.uid in candidate_uids:
            candidates.setdefault(str(cargo.transfer_id or "").strip(), []).append(cargo)
        if cargo.target_order_uid in fallback_orders and not cargo.archive:
            order_cargos.setdefault(cargo.target_order_uid, []).append(cargo)
    return orders, candidates, order_candidates, order_cargos


def _match_snapshot(number, sources):
    orders, candidates, order_candidates, order_cargos = sources
    cargos = candidates.get(number, [])
    if cargos:
        if len(cargos) != 1:
            return None, _issue("ambiguous_original", "系统存在多个相同原单号，不能自动选择")
        cargo = cargos[0]
        order = orders.get(cargo.target_order_uid)
        if not order:
            return None, _issue("order_not_found", "货物尚未归属有效订单")
        if cargo.archive or order.archive:
            return None, _issue("archived_source", "货物或订单已归档，请核对历史归属")
        try:
            return _snapshot(cargo, order), None
        except (InvalidOperation, ValueError):
            return None, _issue("invalid_system_weight", "系统重量无效，请先核对货物资料")
    matches = order_candidates.get(number, [])
    if not matches:
        return None, _issue("original_not_found", "系统找不到该原单号或订单号")
    if len(matches) != 1:
        return None, _issue("ambiguous_order", "系统存在多个相同订单号，不能自动选择")
    order = matches[0]
    if order.archive:
        return None, _issue("archived_source", "订单已归档，请核对历史归属")
    try:
        return _order_snapshot(order, order_cargos.get(order.uid, [])), None
    except (InvalidOperation, ValueError):
        return None, _issue("invalid_system_weight", "系统订单总实重无效，请先核对未归档包裹资料")


def _shipments(snapshot):
    if snapshot.get("match_type") == "order":
        return frozenset((member["cargo_uid"], member["waybill_number"]) for member in snapshot["cargos"])
    return frozenset([(snapshot["cargo_uid"], snapshot["waybill_number"])])


def _prior_source_index(rows):
    shipments = {}
    for row in rows:
        if not row.match_snapshot_json:
            continue
        snapshot = json.loads(row.match_snapshot_json)
        for shipment in _shipments(snapshot):
            shipments.setdefault(shipment, row.batch_id)
    return shipments


def _previous_duplicate(snapshot, index):
    return next((index[shipment] for shipment in sorted(_shipments(snapshot)) if shipment in index), None)


def _prior_sources(session, batch, order_uids, cargo_uids=()):
    sorted_orders, sorted_cargos = sorted(order_uids), sorted(cargo_uids)
    if not sorted_orders and not sorted_cargos:
        return []
    query = session.query(SupplierReconciliationRow).join(SupplierReconciliationBatch).join(
        FinanceDocument, FinanceDocument.id == SupplierReconciliationRow.finance_document_id).filter(
        SupplierReconciliationBatch.vendor_id == batch.vendor_id,
        SupplierReconciliationRow.batch_id != batch.id,
        FinanceDocument.status.in_(["draft", "posted"]))
    # Order-level rows have no single cargo_uid index. Their shipment membership
    # must also be checked when a billed cargo has since moved to another order.
    prior = query.filter(SupplierReconciliationRow.cargo_uid.is_(None)).all()
    for start in range(0, max(len(sorted_orders), len(sorted_cargos)), 500):
        prior.extend(query.filter(SupplierReconciliationRow.cargo_uid.isnot(None),
            or_(SupplierReconciliationRow.order_uid.in_(sorted_orders[start:start + 500]),
                SupplierReconciliationRow.cargo_uid.in_(sorted_cargos[start:start + 500]))).all())
    return prior


def _reconcile_rows(session, batch, rows):
    counts = Counter(row.original_number for row in batch.rows if row.original_number)
    sources = _matching_sources(session, [row.original_number for row in rows if row.original_number])
    prior_index = _prior_source_index(_prior_sources(session, batch, sources[0],
        {cargo.uid for candidates in sources[1].values() for cargo in candidates}
        | {cargo.uid for members in sources[3].values() for cargo in members}))
    for row in rows:
        issues = json.loads(row.parse_issues_json)
        row.status = "exception"
        row.cargo_uid = row.order_uid = row.order_number = row.system_weight = row.system_waybill = row.weight_difference = row.match_snapshot_json = row.fingerprint = None
        if row.original_number and counts[row.original_number] > 1:
            issues.append(_issue("duplicate_in_file", "同一文件原单号重复，请确认一行对应一票最终费用"))
        if not issues:
            snapshot, issue = _match_snapshot(row.original_number, sources)
            if issue:
                issues.append(issue)
            else:
                row.cargo_uid, row.order_uid, row.order_number = snapshot["cargo_uid"], snapshot["order_uid"], snapshot["order_number"]
                row.system_weight, row.system_waybill = Decimal(snapshot["weight"]), snapshot["waybill_number"]
                exact_difference = row.weight - Decimal(snapshot["weight_raw"])
                row.weight_difference = exact_difference.quantize(WEIGHT, rounding=ROUND_HALF_UP)
                row.match_snapshot_json = _json(snapshot)
                order_match = snapshot.get("match_type") == "order"
                if not order_match and row.waybill_number and row.waybill_number != row.system_waybill:
                    issues.append(_issue("waybill_mismatch", "转单号与系统记录不一致"))
                weight_issue = _order_weight_issue(exact_difference) if order_match else _weight_issue(exact_difference)
                if weight_issue:
                    issues.append(weight_issue)
                # Amount/weight changes do not create another final shipment fee.
                row.fingerprint = hashlib.sha256(_json([batch.vendor_id, snapshot["order_uid"],
                    sorted(_shipments(snapshot)), str(row.weight), str(row.final_fee)]).encode()).hexdigest()
                prior = _previous_duplicate(snapshot, prior_index)
                if prior:
                    issues.append(_issue("previous_import_duplicate", f"相同供应商和发运货物已出现在对账批次 {prior}，请核对重复收费或补差"))
        row.issues_json = _json(issues)
        if not issues:
            row.status = "passed"
    # Index membership once; large imports must not scan every row for every row.
    by_shipment = {}
    for other in batch.rows:
        if not other.match_snapshot_json:
            continue
        snapshot = json.loads(other.match_snapshot_json)
        order_match = snapshot.get("match_type") == "order"
        shipments = _shipments(snapshot)
        for shipment in shipments:
            by_shipment.setdefault((shipment, order_match), set()).add(other.id)
    # A whole-order fee and its package fee cannot both enter the same draft.
    for row in rows:
        if row.status != "passed" or not row.match_snapshot_json:
            continue
        snapshot = json.loads(row.match_snapshot_json)
        order_match = snapshot.get("match_type") == "order"
        shipments = _shipments(snapshot)
        overlap = any(by_shipment.get((shipment, not order_match)) for shipment in shipments)
        if order_match:
            overlap = overlap or any(by_shipment.get((shipment, True), set()) - {row.id} for shipment in shipments)
        if overlap:
            row.status = "exception"
            row.issues_json = _json(json.loads(row.issues_json) + [_issue("overlapping_in_file",
                "同一文件的订单费用与包裹费用覆盖相同发运货物，请选择一种对账方式")])


def _create_drafts(session, batch, rows, actor):
    passed = [row for row in rows if row.status == "passed" and row.finance_document_id is None]
    actor_id, actor_name = _actor(actor)
    for start in range(0, len(passed), 200):
        group = passed[start:start + 200]
        body = {"document_type": "vendor_bill", "partner_id": batch.vendor_id,
            "document_date": ledger._today(), "reference": batch.number,
            "memo": f"供应商对账 {batch.period}；仅包含通过行，异常行留在批次 {batch.number}，最终费用未做供应商价格校验"}
        values = ledger._document_values(session, body)
        lines = ledger._lines(session, [{"description": row.description or f"运费：{row.original_number}",
            "quantity": "1", "unit_price": str(row.final_fee), "tax_rate": "0", "order_uid": row.order_uid} for row in group])
        now = ledger._now()
        document = FinanceDocument(**values, number=f"BILL-{uuid.uuid4().hex[:16].upper()}", status="draft", version=1,
            created_by_id=actor_id, created_by_name=actor_name, created_at=now, updated_at=now)
        ledger._set_lines(document, lines)
        session.add(document)
        session.flush()
        for row, line in zip(group, document.lines):
            row.finance_document_id, row.finance_document_line_id = document.id, line.id
        session.flush()
        ledger._audit(session, actor, "finance_document", document.id, "supplier_import_draft",
            after=ledger.document_dict(session, document), reason=f"来源批次 {batch.number}，通过行金额，不包含异常行")


def _refresh_summary(batch):
    batch.total_rows = len(batch.rows)
    batch.passed_rows = sum(row.status == "passed" for row in batch.rows)
    batch.exception_rows = batch.total_rows - batch.passed_rows
    batch.passed_amount = sum((row.final_fee for row in batch.rows if row.status == "passed"), ledger.ZERO)
    batch.total_valid_amount = sum((row.final_fee for row in batch.rows if row.final_fee is not None and row.final_fee > 0), ledger.ZERO)
    batch.status = "exceptions" if batch.exception_rows else "reconciled"


def upload_batch(content, filename, vendor_id, period, actor):
    _actor(actor)
    vendor_id = ledger._id(vendor_id, "供应商")
    if not isinstance(period, str) or not re.fullmatch(r"\d{4}-(0[1-9]|1[0-2])", period):
        raise ValueError("账期必须为 YYYY-MM")
    if not isinstance(filename, str) or not filename.lower().endswith(".xlsx"):
        raise ValueError("仅支持统一模板 .xlsx 文件")
    filename = filename.replace("\\", "/").rsplit("/", 1)[-1]
    if len(filename) > 256:
        raise ValueError("文件名不能超过 256 字")
    parsed = parse_workbook(content)
    digest = hashlib.sha256(content).hexdigest()
    body = {"vendor_id": vendor_id, "period": period, "file_hash": digest}
    def change(session):
        vendor = ledger._partner(session, "vendor", vendor_id, active=True, locked=True)
        existing = session.query(SupplierReconciliationBatch).filter_by(vendor_id=vendor_id, period=period, file_hash=digest).first()
        if existing:
            raise ledger.FinanceConflictError(f"该文件已上传，请查看对账批次 {existing.number}")
        actor_id, actor_name = _actor(actor)
        now = ledger._now()
        batch = SupplierReconciliationBatch(number=f"REC-{uuid.uuid4().hex[:16].upper()}", vendor_id=vendor.id,
            vendor_name=vendor.name, period=period, filename=filename, file_hash=digest, file_bytes=content,
            status="exceptions", version=1, created_by_id=actor_id, created_by_name=actor_name, created_at=now, updated_at=now)
        batch.rows = [SupplierReconciliationRow(**value, status="exception", issues_json="[]") for value in parsed]
        session.add(batch)
        session.flush()
        _reconcile_rows(session, batch, batch.rows)
        _create_drafts(session, batch, batch.rows, actor)
        _refresh_summary(batch)
        session.flush()
        ledger._audit(session, actor, "supplier_reconciliation", batch.id, "upload", after=_batch_dict(session, batch))
        return _detail(session, batch)
    return ledger._run(actor, "supplier_reconciliation.upload", body, change, authorize=_actor)


def reconcile_batch(batch_id, body, actor):
    def change(session):
        batch = session.query(SupplierReconciliationBatch).filter_by(id=ledger._id(batch_id)).populate_existing().with_for_update().first()
        if batch is None:
            raise LookupError("对账批次不存在")
        ledger._version(batch, body)
        ledger._partner(session, "vendor", batch.vendor_id, active=True, locked=True)
        before = _batch_dict(session, batch)
        rows = [row for row in batch.rows if row.status == "exception" and row.finance_document_id is None]
        _reconcile_rows(session, batch, rows)
        _create_drafts(session, batch, rows, actor)
        _refresh_summary(batch)
        batch.version += 1
        batch.updated_at = ledger._now()
        session.flush()
        ledger._audit(session, actor, "supplier_reconciliation", batch.id, "reconcile", before=before, after=_batch_dict(session, batch))
        return _detail(session, batch)
    return ledger._run(actor, f"supplier_reconciliation.{batch_id}.reconcile", body, change, authorize=_actor)


def list_vendors(actor, *, page=1, page_size=30, query_text=""):
    _actor(actor)
    page, page_size = max(int(page or 1), 1), min(max(int(page_size or 30), 1), 200)
    with ledger.SessionFactory() as session:
        query = session.query(FinanceVendor).filter_by(active=True)
        if query_text:
            query = query.filter(FinanceVendor.name.contains(query_text, autoescape=True))
        total = query.count()
        rows = query.order_by(FinanceVendor.name, FinanceVendor.id).offset((page - 1) * page_size).limit(page_size).all()
        return {"items": [{"id": row.id, "name": row.name} for row in rows], "total": total, "page": page, "page_size": page_size}


def list_batches(actor, *, page=1, page_size=30, query_text="", vendor_id=None, period=""):
    _actor(actor)
    page, page_size = max(int(page or 1), 1), min(max(int(page_size or 30), 1), 200)
    with ledger.SessionFactory() as session:
        query = session.query(SupplierReconciliationBatch)
        if vendor_id:
            query = query.filter_by(vendor_id=ledger._id(vendor_id, "供应商"))
        if period:
            query = query.filter_by(period=period)
        if query_text:
            query = query.filter(or_(SupplierReconciliationBatch.number.contains(query_text, autoescape=True),
                SupplierReconciliationBatch.filename.contains(query_text, autoescape=True),
                SupplierReconciliationBatch.vendor_name.contains(query_text, autoescape=True)))
        total = query.count()
        batches = query.order_by(SupplierReconciliationBatch.id.desc()).offset((page - 1) * page_size).limit(page_size).all()
        return {"items": [_batch_dict(session, batch) for batch in batches], "total": total, "page": page, "page_size": page_size}


def get_batch(batch_id, actor, **filters):
    _actor(actor)
    with ledger.SessionFactory() as session:
        batch = session.get(SupplierReconciliationBatch, ledger._id(batch_id))
        if batch is None:
            raise LookupError("对账批次不存在")
        return _detail(session, batch, **filters)


def is_import_document(session, document_id):
    return session.query(SupplierReconciliationRow.id).filter_by(finance_document_id=document_id).first() is not None


def validate_import_document_for_post(session, document):
    """Called inside the ledger posting transaction before any financial effect."""
    if document.document_type != "vendor_bill":
        return
    rows = session.query(SupplierReconciliationRow).filter_by(finance_document_id=document.id).with_for_update().all()
    if not rows:
        return
    lines = {line.id: line for line in document.lines}
    if len(rows) != len(lines) or document.currency != "CNY" or document.is_opening:
        raise ledger.FinanceConflictError("导入应付草稿的明细或币种已变化，请取消后重新上传核对")
    sources = _matching_sources(session, [row.original_number for row in rows])
    for row in rows:
        batch = session.get(SupplierReconciliationBatch, row.batch_id)
        line = lines.get(row.finance_document_line_id)
        if row.status != "passed" or document.partner_type != "vendor" or document.partner_id != batch.vendor_id or line is None:
            raise ledger.FinanceConflictError("导入应付草稿的供应商或来源明细已变化，请重新核对")
        snapshot, issue = _match_snapshot(row.original_number, sources)
        if issue or _json(snapshot) != row.match_snapshot_json:
            raise ledger.FinanceConflictError(f"原单号 {row.original_number} 的单号、重量或订单归属已变化，请重新核对后生成草稿")
        order_match = snapshot.get("match_type") == "order"
        difference = row.weight - Decimal(snapshot["weight_raw"])
        if ((_order_weight_issue(difference) if order_match else _weight_issue(difference)) is not None
            or (not order_match and row.waybill_number and row.waybill_number != snapshot["waybill_number"])):
            raise ledger.FinanceConflictError(f"原单号 {row.original_number} 的当前资料未通过对账校验")
        if (line.order_uid != row.order_uid or line.order_number != row.order_number or line.quantity != Decimal("1")
            or line.unit_price != row.final_fee or line.subtotal != row.final_fee or line.total_amount != row.final_fee
            or line.tax_rate != 0 or line.tax_amount != 0):
            raise ledger.FinanceConflictError("导入应付草稿的订单、金额或税额已变化，请取消后重新上传核对")
    expected = sum((row.final_fee for row in rows), ledger.ZERO)
    if document.total_amount != expected or document.subtotal != expected or document.tax_amount != 0:
        raise ledger.FinanceConflictError("导入应付草稿总额与通过行金额不一致")


def order_costs(order_uid, actor):
    _actor(actor)
    order_uid = ledger._text(order_uid, "订单编号", 32, required=True)
    with ledger.SessionFactory() as session:
        original = aliased(FinanceDocument)
        query = session.query(FinanceDocumentLine, FinanceDocument, SupplierReconciliationRow).join(
            FinanceDocument, FinanceDocument.id == FinanceDocumentLine.document_id).outerjoin(
            original, original.id == FinanceDocument.original_document_id).outerjoin(
            SupplierReconciliationRow, SupplierReconciliationRow.finance_document_line_id == FinanceDocumentLine.id).filter(
            FinanceDocumentLine.order_uid == order_uid, FinanceDocument.document_type.in_(["vendor_bill", "vendor_credit"]),
            FinanceDocument.is_opening == False,  # noqa: E712
            or_(original.id.is_(None), original.is_opening == False),  # noqa: E712
            FinanceDocument.status.in_(["posted", "draft"]))
        rows = query.order_by(FinanceDocument.document_date.desc(), FinanceDocument.id.desc(), FinanceDocumentLine.sequence).all()
        if not session.get(OrderTable, order_uid) and not rows:
            raise LookupError("订单不存在")
        confirmed = pending = net = ledger.ZERO
        items = []
        for line, document, source in rows:
            sign = -1 if document.document_type == "vendor_credit" else 1
            total, subtotal = line.total_amount * sign, line.subtotal * sign
            if document.status == "posted":
                confirmed += total
                net += subtotal
            else:
                pending += total
            if len(items) < 100:
                items.append({"document_id": document.id, "number": document.number, "document_type": document.document_type,
                    "status": document.status, "line_id": line.id, "description": line.description,
                    "total_amount": ledger._str(total), "subtotal": ledger._str(subtotal), "currency": document.currency,
                    "source_batch_id": source.batch_id if source else None, "source_row": source.source_row if source else None,
                    "match_type": ("cargo" if source.cargo_uid else "order" if source.order_uid else None) if source else None,
                    "source_original_number": source.original_number if source else None,
                    "source_system_weight": ledger._str(source.system_weight, WEIGHT) if source and source.system_weight is not None else None})
        return {"confirmed_amount": ledger._str(confirmed), "pending_amount": ledger._str(pending),
            "confirmed_net_amount": ledger._str(net), "items": items, "total": len(rows),
            "scope": "已生效及草稿供应商账单（含手工账单）；贷项以负数冲减，汇总包含全部明细，列表最多显示 100 条"}
