from io import BytesIO

from flask import Blueprint, jsonify, request, send_file
from openpyxl import load_workbook
import json

from Auth import verify, allow, get_token_payload
from data import (
    order_model,
    get_customer_record,
    list_customer_records,
    list_receiver_records,
    del_receiver_record,
    add_customer_record,
)
from data.customer.Customer import (
    del_customer_record,
    get_customer_detail,
    get_all_customers_detail,
    update_customer_info,
    ALLOWED_ROLES,
)
from services.order_edit_policy import OrderEditConflictError
from utils.audit_log import (
    log_sensitive_access,
    ACTION_VIEW_ALL,
    ACTION_VIEW_SINGLE,
    ACTION_UPDATE,
    ACTION_CREATE,
    ACTION_DELETE,
)

customer_bp = Blueprint("customer", __name__, url_prefix="/api")


# 客户列表
@customer_bp.get("/customer_list")
@verify
@allow()
def get_customer_list():
    return jsonify({"code": 200, "data": get_customer_record()})


# 客户名下的收件人列表（仅姓名）
@customer_bp.get("/receiver_list/<int:customer_id>")
@verify
@allow()
def get_receiver_list(customer_id):
    return jsonify({"code": 200, "data": get_customer_record(customer_id=customer_id)})


# 客户名下收件人的完整信息列表
@customer_bp.get("/receiver_info_list/<int:customer_id>")
@verify
@allow()
def get_receiver_info_list(customer_id):
    receivers = get_customer_record(customer_id=customer_id)
    ids = [r["receiver_id"] for r in receivers]
    return jsonify({"code": 200, "data": list_receiver_records(ids)})


# 单个收件人完整信息
@customer_bp.get("/receiver_info/<receiver_id>")
@verify
@allow()
def get_receiver_info(receiver_id):
    return jsonify(
        {"code": 200, "data": list_customer_records(receiver_id=receiver_id)}
    )


# 删除收件人
@customer_bp.post("/del/<int:customer_id>/<receiver_id>")
@verify
@allow()
def del_receiver(customer_id, receiver_id):
    del_receiver_record(customer_id=customer_id, receiver_id=receiver_id)
    return jsonify({"code": 200, "msg": "success"})


# 删除客户（级联删除所有收件人）
@customer_bp.post("/del_customer/<int:customer_id>")
@verify
@allow(role=ALLOWED_ROLES)
def del_customer(customer_id):
    """删除客户及其所有收件人（敏感操作，需要权限）"""
    payload = get_token_payload()

    # 记录审计日志
    log_sensitive_access(
        user_id=payload.get("id"),
        user_role=payload.get("role"),
        user_name=payload.get("name"),
        action=ACTION_DELETE,
        customer_name=f"ID:{customer_id}",
        ip_address=request.remote_addr,
    )

    try:
        del_customer_record(customer_id=customer_id)
    except OrderEditConflictError as exc:
        return jsonify({"code": 409, "msg": str(exc)}), 409
    return jsonify({"code": 200, "msg": "success"})


# 创建客户
@customer_bp.post("/customer")
@verify
@allow(role=ALLOWED_ROLES)
def create_customer():
    """创建客户（仅客户基本信息，敏感操作，需要权限）。

    Request body:
    {
        "name": "客户名称",    # 必填
        "phone": "客户电话",   # 可选
        "wechat": "客户微信"   # 可选
    }

    Returns:
        customer_id: 客户ID
    """
    from data.customer.Customer import CustomerTable
    from data.database import SessionFactory

    payload = get_token_payload()
    record = json.loads(request.data)
    name = record.get("name", "").strip()

    if not name:
        return jsonify({"code": 400, "msg": "name is required"}), 400

    phone = record.get("phone", "").strip() or None
    wechat = record.get("wechat", "").strip() or None

    # 记录审计日志
    log_sensitive_access(
        user_id=payload.get("id"),
        user_role=payload.get("role"),
        user_name=payload.get("name"),
        action=ACTION_CREATE,
        customer_name=name,
        ip_address=request.remote_addr,
    )

    s = SessionFactory()
    try:
        # 检查是否已存在
        existing = s.query(CustomerTable).filter_by(name=name).first()
        if existing:
            return (
                jsonify(
                    {
                        "code": 400,
                        "msg": "Customer already exists",
                        "data": {"customer_id": existing.id},
                    }
                ),
                400,
            )

        customer = CustomerTable(name=name, phone=phone, wechat=wechat)
        s.add(customer)
        s.commit()
        customer_id = customer.id
        return jsonify(
            {"code": 200, "msg": "success", "data": {"customer_id": customer_id}}
        )
    except Exception as e:
        s.rollback()
        return (
            jsonify({"code": 500, "msg": f"Failed to create customer: {str(e)}"}),
            500,
        )
    finally:
        s.close()


# 下载预填地址的 Excel 模板
@customer_bp.get("/excel")
def get_excel_template():
    wb = load_workbook("static/template.xlsx")
    ws = wb.active
    ws["A17"] = request.args.get("address", "")
    ws["A19"] = request.args.get("zip_code", "")
    ws["C19"] = request.args.get("country", "")
    ws["A20"] = request.args.get("receiver", "")
    ws["C20"] = request.args.get("receiver_phone", "")
    stream = BytesIO()
    wb.save(stream)
    stream.seek(0)
    return send_file(
        stream,
        as_attachment=True,
        mimetype="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        download_name="",
    )


# ---------------------------------------------------------------------------
# 客户敏感信息管理接口（phone, wechat）
# ---------------------------------------------------------------------------


# 获取所有客户详细信息（包括敏感字段，需要权限）
@customer_bp.get("/customers_detail")
@verify
@allow(role=ALLOWED_ROLES)
def get_customers_detail_list():
    """获取所有客户的详细信息，包括phone和wechat（需要权限）"""
    payload = get_token_payload()

    # 记录审计日志
    log_sensitive_access(
        user_id=payload.get("id"),
        user_role=payload.get("role"),
        user_name=payload.get("name"),
        action=ACTION_VIEW_ALL,
        ip_address=request.remote_addr,
    )

    customers = get_all_customers_detail(include_sensitive=True)
    return jsonify({"code": 200, "data": customers})


# 获取单个客户详细信息（包括敏感字段，需要权限）
@customer_bp.get("/customer_detail/<int:customer_id>")
@verify
@allow(role=ALLOWED_ROLES)
def get_customer_detail_info(customer_id):
    """获取单个客户的详细信息，包括phone和wechat（需要权限）"""
    payload = get_token_payload()

    # 记录审计日志
    log_sensitive_access(
        user_id=payload.get("id"),
        user_role=payload.get("role"),
        user_name=payload.get("name"),
        action=ACTION_VIEW_SINGLE,
        customer_name=f"ID:{customer_id}",
        ip_address=request.remote_addr,
    )

    customer = get_customer_detail(customer_id=customer_id, include_sensitive=True)
    if not customer:
        return jsonify({"code": 404, "msg": "Customer not found"}), 404
    return jsonify({"code": 200, "data": customer})


# 修改客户信息
@customer_bp.put("/customer/<int:customer_id>")
@verify
@allow(role=ALLOWED_ROLES)
def update_customer(customer_id):
    """修改客户基本信息。

    Request body:
    {
        "name": "新客户名称",   # 可选
        "phone": "客户电话",    # 可选
        "wechat": "客户微信"    # 可选
    }
    """
    from data.customer.Customer import CustomerTable
    from data.database import SessionFactory

    record = json.loads(request.data)
    payload = get_token_payload()

    # 记录审计日志
    log_sensitive_access(
        user_id=payload.get("id"),
        user_role=payload.get("role"),
        user_name=payload.get("name"),
        action=ACTION_UPDATE,
        customer_name=f"ID:{customer_id}",
        ip_address=request.remote_addr,
    )

    s = SessionFactory()
    try:
        customer = s.query(CustomerTable).filter_by(id=customer_id).first()
        if customer is None:
            return jsonify({"code": 404, "msg": "Customer not found"}), 404

        # 更新字段
        if "name" in record:
            new_name = record["name"].strip()
            if new_name and new_name != customer.name:
                # 检查名称是否与其他客户冲突
                existing = (
                    s.query(CustomerTable)
                    .filter(
                        CustomerTable.name == new_name, CustomerTable.id != customer_id
                    )
                    .first()
                )
                if existing:
                    return (
                        jsonify({"code": 400, "msg": "Customer name already exists"}),
                        400,
                    )
                from data.DB import OrderTable
                from data.order_version import (
                    new_correlation_id,
                    record_mutation,
                    snapshot_order,
                )

                old_name = customer.name
                order_uids = sorted(order.uid for order in customer.orders)
                if order_uids:
                    orders = (
                        s.query(OrderTable)
                        .filter(OrderTable.uid.in_(order_uids))
                        .order_by(OrderTable.uid.asc())
                        .with_for_update()
                        .all()
                    )
                else:
                    orders = []
                customer = (
                    s.query(CustomerTable)
                    .filter(CustomerTable.id == customer_id)
                    .populate_existing()
                    .with_for_update()
                    .first()
                )
                if customer is None or customer.name != old_name:
                    raise OrderEditConflictError("客户资料已变化，请刷新后重试")
                before_by_order = {
                    order_uid: snapshot_order(s, order_uid)
                    for order_uid in order_uids
                }
                customer.name = new_name
                for order in orders:
                    order.customer_name = new_name
                s.flush()
                correlation_id = new_correlation_id()
                for order_uid in order_uids:
                    record_mutation(
                        s,
                        order_uid,
                        before=before_by_order[order_uid],
                        after=snapshot_order(s, order_uid),
                        action="customer_updated",
                        summary="修改客户名称",
                        correlation_id=correlation_id,
                    )

        if "phone" in record:
            customer.phone = record["phone"].strip() if record["phone"] else None

        if "wechat" in record:
            customer.wechat = record["wechat"].strip() if record["wechat"] else None

        s.commit()
        return jsonify(
            {
                "code": 200,
                "msg": "success",
                "data": {
                    "id": customer.id,
                    "name": customer.name,
                    "phone": customer.phone,
                    "wechat": customer.wechat,
                },
            }
        )
    except OrderEditConflictError as e:
        s.rollback()
        return jsonify({"code": 409, "msg": str(e)}), 409
    except Exception as e:
        s.rollback()
        return (
            jsonify({"code": 500, "msg": f"Failed to update customer: {str(e)}"}),
            500,
        )
    finally:
        s.close()
